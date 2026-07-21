"""補助シートの自動差し込み（宅建業者追記欄・付帯設備表・物件状況等報告書）.

セル座標について:
- 宅建業者追記欄は FRK 固定様式のため、ラベル位置は安定。表紙欄と同じ論理項目を
  実測した座標に割り当てる（36-1/37-1 共通レイアウト）。
- 付帯設備表・物件状況等報告書は項目数が多く反復的なため、**テンプレートを実行時に
  走査**してラベル（「発見していない」「売主」「有」）直左の □ を特定する。
  座標のハードコードを避け、様式改訂に強くする。

安全方針（重要）:
- 追記欄（＝会社の確定情報）は常時適用してよい。
- 付帯設備表・物件状況等報告書の既定値は **事実の断定**（例「発見していない」）を
  含むため、既定では適用しない。deal フラグで明示 opt-in した時のみ適用し、
  必ず「未検証の初期値・現地確認必須」の警告を出す。
"""

from __future__ import annotations

from typing import Any

import house_style
from cellmaps import ON, _split_menkyo, _split_tel, _split_toroku

SHEET_TEKKI = "宅建業者追記欄"

# ── 宅建業者追記欄の座標（左=売主側 / 右=媒介側）─────────────────
_TEKKI_L = {
    "taiyo_urinushi": "T6",                          # ■売主
    "menkyo_ken": "H8", "menkyo_kai": "O8", "menkyo_no": "S8",
    "shozai": "H10", "shozai2": "H12", "tel": ("H14", "N14", "T14"),
    "shomei": "H16", "daihyo": "H18",
    "toroku_ken": "H20", "toroku_no": "R20", "shimei": "H22",
    "jimusho": "H24", "jimusho_shozai": "H26", "jimusho_shozai2": "H28",
    "ts_tel": ("H30", "N30", "T30"),
    "member": "C32", "kyokai": "J34", "kyokai_addr": "J36",
    "honbu": "J38", "honbu_addr": "J40",
    "bensai": "J42", "bensai_addr": "J44",
}
_TEKKI_R = {
    "taiyo_baikai": "AH6",                            # ■媒介
    "menkyo_ken": "AF8", "menkyo_kai": "AM8", "menkyo_no": "AQ8",
    "shozai": "AF10", "shozai2": "AF12", "tel": ("AF14", "AL14", "AR14"),
    "shomei": "AF16", "daihyo": "AF18",
    "toroku_ken": "AF20", "toroku_no": "AP20", "shimei": "AF22",
    "jimusho": "AF24", "jimusho_shozai": "AF26", "jimusho_shozai2": "AF28",
    "ts_tel": ("AF30", "AL30", "AR30"),
    "member": "AA32", "kyokai": "AH34", "kyokai_addr": "AH36",
    "honbu": "AH38", "honbu_addr": "AH40",
    "bensai": "AH42", "bensai_addr": "AH44",
}


def _seller_b_gyosha() -> dict[str, Any]:
    """売主側の既定業者（株式会社Martial Arts）。"""
    return dict(house_style.SELLER_B_MASTER)


def _seller_b_torikiishi() -> dict[str, Any]:
    """売主側の既定取引士。

    実務指定に従い、追記欄には**指定の説明宅建士（先頭＝小玉 浩之）**を入れる。
    退職期日の注意は bc_service._torikiishi_warnings が別途 warning で出すため、
    ここでは指定どおりの取引士を返す（案件マスタ bc_torikiishi_* があれば
    呼び出し側でそちらが優先される）。
    """
    cand = house_style.SELLER_B_TORIKIISHI
    return dict(cand[0]) if cand else {}


def _default_baikai() -> tuple[dict[str, Any], dict[str, Any]]:
    """媒介側（右欄）の既定業者と取引士を (gyosha, torikiishi) で返す。

    実務指定に従い、案件マスタに媒介業者(bc_baikai_gyosha_*)が無い場合の既定として
    先頭の媒介業者マスタ（東洋建設ホーム）を充てる。案件で客付業者が異なる場合は
    bc_baikai_gyosha_* で上書きされる（呼び出し側で優先）。
    """
    m = house_style.BAIKAI_GYOSHA_MASTER[0] if house_style.BAIKAI_GYOSHA_MASTER else {}
    gyosha = {
        "menkyo_no": m.get("menkyo_no"),
        "shozai": m.get("shozai"),
        "tel": m.get("tel"),
        "shomei": m.get("shomei"),
        "daihyo": m.get("daihyo"),
    }
    tori = {
        "shimei": m.get("torikiishi_shimei"),
        "toroku_no": m.get("torikiishi_toroku_no"),
    }
    return gyosha, tori


def _get(o: Any, k: str) -> Any:
    if o is None:
        return None
    if isinstance(o, dict):
        return o.get(k)
    return getattr(o, k, None)


def _block_values(coords: dict, gyosha: Any, tori: Any,
                  taiyo_key: str) -> dict[str, Any]:
    """1ブロック分の {coord: 値} を返す（値なし項目は None＝クリア）。"""
    mk, mkai, mno = _split_menkyo(_get(gyosha, "menkyo_no"))
    t1, t2, t3 = _split_tel(_get(gyosha, "tel"))
    rk, rno = _split_toroku(_get(tori, "toroku_no"))
    s1, s2, s3 = _split_tel(_get(tori, "tel"))
    member = _get(gyosha, "is_kyokai_member")
    out: dict[str, Any] = {
        coords["menkyo_ken"]: mk, coords["menkyo_kai"]: mkai, coords["menkyo_no"]: mno,
        coords["shozai"]: _get(gyosha, "shozai"), coords["shozai2"]: None,
        coords["shomei"]: _get(gyosha, "shomei"), coords["daihyo"]: _get(gyosha, "daihyo"),
        coords["toroku_ken"]: rk, coords["toroku_no"]: rno,
        coords["shimei"]: _get(tori, "shimei"),
        coords["jimusho"]: _get(tori, "jimusho"),
        coords["jimusho_shozai"]: _get(tori, "jimusho_shozai"),
        coords["jimusho_shozai2"]: None,
        coords["member"]: (ON if member else None),
        coords["kyokai"]: _get(gyosha, "hosho_kyokai"),
        coords["kyokai_addr"]: _get(gyosha, "hosho_kyokai_addr"),
        coords["honbu"]: _get(gyosha, "hosho_honbu"),
        coords["honbu_addr"]: _get(gyosha, "hosho_honbu_addr"),
        coords["bensai"]: _get(gyosha, "bensai_kyotaku"),
        coords["bensai_addr"]: _get(gyosha, "bensai_kyotaku_addr"),
    }
    if taiyo_key in coords:
        out[coords[taiyo_key]] = ON        # 取引態様チェック
    for coord, val in zip(coords["tel"], (t1, t2, t3)):
        out[coord] = val
    for coord, val in zip(coords["ts_tel"], (s1, s2, s3)):
        out[coord] = val if s1 is not None else None
    return out


def tekki_values(bc: Any) -> dict[str, dict[str, Any]]:
    """宅建業者追記欄シートの {シート名: {coord: 値}} を返す。

    左＝売主側（bc.gyosha or 既定 Martial Arts）、右＝媒介側（bc.baikai_gyosha）。
    """
    left_g = _get(bc, "gyosha") or _seller_b_gyosha()
    left_t = _get(bc, "torikiishi") or _seller_b_torikiishi()
    right_g = _get(bc, "baikai_gyosha")
    right_t = _get(bc, "baikai_torikiishi")
    if not right_g and not right_t:
        # 案件マスタに媒介業者指定が無ければ既定（東洋建設ホーム）を右欄に充てる。
        right_g, right_t = _default_baikai()

    vals = _block_values(_TEKKI_L, left_g, left_t, "taiyo_urinushi")
    if right_g or right_t:
        vals.update(_block_values(_TEKKI_R, right_g, right_t, "taiyo_baikai"))
    return {SHEET_TEKKI: vals}


# ── 付帯設備表・物件状況報告書の既定（走査で座標を導出）──────────
def _left_checkbox(ws: Any, row: int, label_col: int) -> str | None:
    """label セルの左側にある直近の □ セルの座標を返す（無ければ None）。"""
    from openpyxl.utils import get_column_letter
    for c in range(label_col - 1, max(label_col - 6, 0), -1):
        v = ws.cell(row, c).value
        if isinstance(v, str) and v.strip() == "□":
            return f"{get_column_letter(c)}{row}"
    return None


def scan_kokuchi_defaults(ws: Any) -> dict[str, str]:
    """物件状況等報告書: 各項目の「□売主」と「□発見していない」に ■ を立てる。"""
    out: dict[str, str] = {}
    for row in range(1, ws.max_row + 1):
        for c in range(1, min(ws.max_column + 1, 40)):
            v = ws.cell(row, c).value
            if not isinstance(v, str):
                continue
            s = v.strip()
            if s == "発見していない" or s == "売主":
                cb = _left_checkbox(ws, row, c)
                if cb:
                    out[cb] = ON
    return out


# 中古戸建の標準装備（既定で「設備有無=有／故障不具合=無」を立てる設備名）。
# 現地確認後に変更する前提。ここに無い設備（食洗機・ディスポーザー等）は既定では
# 有無を断定せず空欄のままにする（存在しない設備を「有」と誤断定しないため）。
_SETSUBI_ARI_KEYS = (
    "給湯設備", "流し台", "混合水栓", "コンロ",
    "シャワー", "トイレ設備", "防水パン", "洗濯用水栓",
)
# 給湯設備の設置箇所（□の直左ラベル）。標準の3箇所を既定で■にする。
_KYUTO_LOCATIONS = ("キッチン", "浴室", "洗面所")


def _setsubi_columns(ws: Any) -> tuple[int | None, int | None, int | None]:
    """ヘッダから「設備の有無」「故障不具合」「故障・不具合『有』の場合」の列を返す。"""
    umu = koshou = koshou_end = None
    for row in range(1, min(ws.max_row + 1, 30)):
        for c in range(1, min(ws.max_column + 1, 55)):
            v = ws.cell(row, c).value
            if not isinstance(v, str):
                continue
            s = v.strip()
            if s == "設備の有無":
                umu = c
            elif s == "故障不具合":
                koshou = c
            elif "故障" in s and "有」の場合" in s:
                koshou_end = c
        if umu and koshou:
            break
    return umu, koshou, koshou_end


def scan_setsubi_defaults(ws: Any) -> dict[str, str]:
    """付帯設備表: 中古戸建の標準装備に「設備有無=有／故障不具合=無」を立てる。

    列を認識し、設備有無欄の □有 と 故障不具合欄の □無 を区別して■にする。
    対象は _SETSUBI_ARI_KEYS の標準装備のみ（存在断定を避けるため全設備は埋めない）。
    給湯設備行では設置箇所（キッチン・浴室・洗面所）の□も■にする。
    現地確認後に変更する前提（呼び出し側が警告を出す）。
    """
    out: dict[str, str] = {}
    umu_col, koshou_col, koshou_end = _setsubi_columns(ws)
    if not (umu_col and koshou_col):
        return out
    koshou_end = koshou_end or (koshou_col + 6)

    for row in range(1, ws.max_row + 1):
        # この行の設備名ラベル（カテゴリ列D=4 と小項目列K=11）を集める。
        labels = []
        for lc in (4, 11):
            v = ws.cell(row, lc).value
            if isinstance(v, str) and v.strip():
                labels.append(v.strip())
        if not any(any(key in lab for lab in labels) for key in _SETSUBI_ARI_KEYS):
            continue

        # 設備有無欄の「有」→左□を■
        for c in range(umu_col, koshou_col):
            v = ws.cell(row, c).value
            if isinstance(v, str) and v.strip() == "有":
                cb = _left_checkbox(ws, row, c)
                if cb:
                    out[cb] = ON
        # 故障不具合欄の「無」→左□を■
        for c in range(koshou_col, koshou_end):
            v = ws.cell(row, c).value
            if isinstance(v, str) and v.strip() == "無":
                cb = _left_checkbox(ws, row, c)
                if cb:
                    out[cb] = ON

        # 給湯設備: 設置箇所（キッチン・浴室・洗面所）の□も既定で■。
        # 設置箇所ラベルは給湯設備行の直下数行に跨る（洗面所は+2行目等）ため広めに走査。
        # 一致は _KYUTO_LOCATIONS の完全一致のみ（「浴室設備」等の別設備は拾わない）。
        if any("給湯設備" in lab for lab in labels):
            for r2 in range(row, min(row + 5, ws.max_row + 1)):
                for c in range(1, min(ws.max_column + 1, 26)):
                    v = ws.cell(r2, c).value
                    if isinstance(v, str) and v.strip() in _KYUTO_LOCATIONS:
                        cb = _left_checkbox(ws, r2, c)
                        if cb:
                            out[cb] = ON
    return out
